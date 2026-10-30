"""Optional Excel export of the replacement mapping (old -> new per domain).

The file contains ORIGINAL client values: it is written only to a user-chosen
path outside the published package (pipeline validates the path) and must be
handled like the .snpanon project. Streaming both ways: rows arrive from
SQLite as an iterator and the workbook is written in write_only mode, so memory
stays bounded. Excel cell and sheet limits are handled explicitly
(UTF-16-safe truncation, control-character cleanup, sheet rollover).
"""
from __future__ import annotations

import re
import json
import sqlite3
from pathlib import Path

from .errors import AnonymizerError

HEADER = ("Домен", "Исходное значение", "Замена")
CELL_LIMIT = 32_000
CELL_MARKER = " ...[truncated]"
CELL_UTF16_LIMIT = 32_767
ROWS_PER_SHEET = 1_000_000
ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\uFFFE\uFFFF]")
LABEL_BYTES = 24 * 1024
LABEL_MARKER = "\n… (список полей сокращён)"
DOMAIN_TITLES = {"LOCATION": "Локация", "PRODUCT": "Продукт", "RESOURCE": "Ресурс",
                 "BOM": "Спецификация", "TRANSPORT_TYPE": "Тип транспорта", "CLIENT": "Клиент"}


def iter_mapping_rows(project, check=lambda: None):
    """Private display labels; durable mapping keys and values stay untouched.

    Merge two disk-sorted streams, retaining only one binding and one bounded
    label. Saved bindings include fields/namespaces absent from this import.
    """
    from .connectivity import _hash, load_profile

    contract = project.contract
    if contract is None:
        yield from project.iter_mapping()
        return
    builtin = contract.profile_hash == _hash(load_profile("builtin:optimizer-v1"))
    mappings, bindings = project.iter_mapping(), None
    interrupted = None

    def progress():
        nonlocal interrupted
        try:
            check()
        except BaseException as error:
            interrupted = error
            return 1
        return 0

    def invalid():
        raise AnonymizerError("Не удалось определить домен карты замен: некорректные сохранённые привязки проекта.")

    def decode(row):
        if (not isinstance(row[1], str) or not isinstance(row[2], str)
                or len(row[1]) > 1024 or len(row[2]) > 2048):
            invalid()
        try:
            node, semantic = json.loads(row[1]), json.loads(row[2])
            namespace = semantic["namespace"]
            if (not isinstance(namespace, str) or not 0 < len(namespace) <= 63
                    or not isinstance(node, list) or len(node) != 3
                    or any(not isinstance(v, str) or not 0 < len(v) <= 63 for v in node)):
                invalid()
        except (ValueError, KeyError, TypeError):
            invalid()
        if contract.profile["namespace_rule"] != "shared" and node[0] != namespace:
            invalid()
        return node, namespace

    project.db.set_progress_handler(progress, 1000)
    try:
        bindings = project.db.execute("SELECT domain,node,semantic FROM connectivity_bindings ORDER BY domain,node")

        def advance():
            row = next(bindings, None)
            if row is not None and (not isinstance(row[0], str) or len(row[0]) > 200):
                invalid()
            return row

        current = advance()
        previous, label = None, None
        for domain, old, new in mappings:
            check()
            if not domain.startswith("connectivity:"):
                yield domain, old, new
                continue
            if domain != previous:
                while current is not None and current[0] < domain:
                    check()
                    current = advance()
                if current is None or current[0] != domain:
                    invalid()
                _, namespace = decode(current)
                spec = next((spec for spec in contract.profile["domains"]
                             if domain == "connectivity:" + contract.profile_hash + ":" + _hash([namespace, spec["id"]])), None)
                if spec is None:
                    invalid()
                ident = spec["id"]
                title = f"{DOMAIN_TITLES[ident]} ({ident})" if builtin and ident in DOMAIN_TITLES else f"Домен {ident}"
                shared = contract.profile["namespace_rule"] == "shared"
                label = title + (", общий namespace:" if shared else f", схема {namespace}:")
                size, truncated = len(label.encode("utf-8")), False
                while current is not None and current[0] == domain:
                    check()
                    node, saved_namespace = decode(current)
                    if saved_namespace != namespace:
                        invalid()
                    text = "\n" + ".".join(node if shared else node[1:])
                    if not truncated:
                        addition = len(text.encode("utf-8"))
                        if size + addition + len(LABEL_MARKER.encode("utf-8")) <= LABEL_BYTES:
                            label += text
                            size += addition
                        else:
                            label += LABEL_MARKER
                            truncated = True
                    current = advance()
                previous = domain
            yield label, old, new
    except sqlite3.OperationalError:
        if interrupted is not None:
            raise interrupted from None
        raise
    finally:
        mappings.close()
        if bindings is not None:
            bindings.close()
        project.db.set_progress_handler(None, 0)


def utf16_units(text):
    return len(text.encode("utf-16-le", "surrogatepass")) // 2


def cell_value(text):
    """Excel rejects control characters and caps cells in UTF-16 units; make any value safe."""
    clean = ILLEGAL.sub("", text)
    if len(clean) <= CELL_LIMIT and utf16_units(clean) <= CELL_UTF16_LIMIT:
        return clean
    cut = clean[:CELL_LIMIT - len(CELL_MARKER)]
    if utf16_units(cut + CELL_MARKER) > CELL_UTF16_LIMIT:
        # Astral characters count as two units; find the largest safe prefix.
        low, high = 0, len(cut)
        while low < high:
            middle = (low + high + 1) // 2
            if utf16_units(cut[:middle] + CELL_MARKER) <= CELL_UTF16_LIMIT:
                low = middle
            else:
                high = middle - 1
        cut = cut[:low]
    return cut + CELL_MARKER


def write_mapping_xlsx(path, rows, check=lambda: None):
    """rows: iterable of (domain, old, new); returns the number of data rows."""
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Alignment
    except ImportError:
        raise AnonymizerError("Библиотека Excel не установлена. Пересоберите приложение или установите openpyxl.") from None
    workbook = Workbook(write_only=True)
    sheets = []
    alignment = Alignment(wrap_text=True, vertical="top")
    try:
        sheet, sheet_number, written, total = None, 0, 0, 0
        for domain, old, new in rows:
            check()
            if sheet is None or written >= ROWS_PER_SHEET:
                if sheet is not None:
                    sheet.auto_filter.ref = f"A1:C{written + 1}"
                sheet_number += 1
                sheet = workbook.create_sheet("mapping" if sheet_number == 1 else f"mapping_{sheet_number}")
                sheets.append(sheet)
                sheet.freeze_panes = "A2"
                for column, width in (("A", 72), ("B", 42), ("C", 42)):
                    sheet.column_dimensions[column].width = width
                sheet.append(HEADER)
                written = 0
            cells = []
            for value in (domain, old, new):
                cell = WriteOnlyCell(sheet, value=cell_value(value))
                # Values are data, including strings that Excel would infer as
                # formulas or errors (e.g. '=1+1' and '#N/A').
                cell.data_type = "s"
                cell.alignment = alignment
                cells.append(cell)
            sheet.append(cells)
            written += 1
            total += 1
        if sheet is None:
            sheet = workbook.create_sheet("mapping")
            sheets.append(sheet)
            sheet.freeze_panes = "A2"
            for column, width in (("A", 72), ("B", 42), ("C", 42)):
                sheet.column_dimensions[column].width = width
            sheet.append(HEADER)
        sheet.auto_filter.ref = f"A1:C{written + 1}"
        workbook.save(str(path))
    except AnonymizerError:
        raise
    except OSError as error:
        raise AnonymizerError(f"Не удалось записать файл карты замен ({type(error).__name__}). "
                              "Проверьте путь и свободное место.") from None
    finally:
        # Aborted write_only sheets keep row-writer generators; GC of an open
        # one raises "I/O operation on closed file" unless closed here.
        for sheet in sheets:
            try:
                sheet.close()
            except Exception:
                pass
            # close() alone leaves write_only worksheet XML on error/cancel.
            # Clean only this export's writers; other workbooks may be active.
            writer = getattr(sheet, "_writer", None)
            if writer is not None:
                try:
                    writer.close()
                except Exception:
                    pass
                try:
                    if isinstance(writer.out, str) and Path(writer.out).is_file():
                        writer.cleanup()
                except (OSError, ValueError):
                    pass  # cleanup must not replace the original failure
        try:
            workbook.close()
        except Exception:
            pass  # a failed save leaves nothing to finalise
    return total
