"""Synthetic generic-mode acceptance; no model names or FK are inferred."""
import csv
import json
import zipfile

import pytest
from openpyxl import Workbook

from anonymizer.errors import AnonymizerError
from anonymizer.pipeline import Options, run
from anonymizer.runlog import RunLog


def output_rows(output):
    manifest = json.loads((output / "csv/manifest.json").read_text(encoding="utf-8"))
    result = {}
    for entry in manifest["tables"]:
        with (output / "csv" / entry["file"]).open(encoding="utf-8", newline="") as stream:
            result[entry["table"]] = list(csv.DictReader(stream))
    return result


@pytest.mark.parametrize("table", ["sku", "custom_unlisted"])
def test_full_bypass_keeps_extra_columns_without_business_keys(tmp_path, table):
    source = tmp_path / "source"
    source.mkdir()
    data = ("note,order_id,qty,day,uom,abc,activeflag,sys_id\n"
            "SYN_SECRET,999,12345678901234567890.123456789,2026-10-29,kg,A,Y,123\n")
    file = source / f"{table}.csv"
    file.write_text(data, encoding="utf-8")
    output = tmp_path / "out"
    report = run(Options(source, output, tmp_path / "private.snpanon", bypass_model=True))
    assert report["rows"] == report["tables"] == 1
    assert not report["model_applied"] and report["structure_mode"] == "automatic_without_model"
    assert not report["relationships_checked"] and not report["keys_checked"]
    assert not report["skipped_tables"] and report["validation"]["preserved_fields_equal"]
    row, = output_rows(output)[table]
    assert set(row) == {"note", "order_id", "qty", "day", "uom", "abc", "activeflag", "sys_id"}
    assert row["sys_id"] == ""
    assert (row["order_id"], row["qty"], row["day"]) == ("999", "12345678901234567890.123456789", "2026-10-29")
    for key, original in (("note", "SYN_SECRET"), ("uom", "kg"), ("abc", "A"), ("activeflag", "Y")):
        assert row[key] != original
    assert file.read_text(encoding="utf-8") == data
    for file in output.rglob("*"):
        assert file.suffix not in {".xlsx", ".snpanon", ".sqlite"}


def test_cli_mixed_zip_known_and_unknown(tmp_path):
    from anonymizer.cli import main
    folder = tmp_path / "source"
    folder.mkdir()
    (folder / "sku.csv").write_text("private_note,amount\nSYN_UNKNOWN,42.125\n", encoding="utf-8")
    book = Workbook()
    book.active.append(["code", "amount"])
    book.active.append(["001", 12.5])
    book.save(folder / "custom.xlsx")
    book.close()
    source = tmp_path / "input.zip"
    with zipfile.ZipFile(source, "w") as archive:
        for file in folder.iterdir():
            archive.write(file, "data/" + file.name)
    for name in ("first", "second"):
        assert main(["--input", str(source), "--output", str(tmp_path / name),
                     "--project", str(tmp_path / "private.snpanon"), "--bypass-model"]) == 0
        rows = output_rows(tmp_path / name)
        assert rows["sku"][0]["amount"] == "42.125"
        assert rows["custom"][0]["amount"] == "12.5"
        assert rows["custom"][0]["code"] != "001"
    assert (tmp_path / "first/dump.sql").read_bytes() == (tmp_path / "second/dump.sql").read_bytes()


def test_standard_mode_and_project_domains_remain_independent(tmp_path):
    folder = tmp_path / "source"
    folder.mkdir()
    (folder / "item.csv").write_text("item,proddescr\nSYN_CODE,SYN_DESCRIPTION\n", encoding="utf-8")
    project = tmp_path / "private.snpanon"
    for name, mode in (("standard", False), ("generic", True), ("standard-again", False), ("generic-again", True)):
        run(Options(folder, tmp_path / name, project, bypass_model=mode))
    assert (tmp_path / "standard/dump.sql").read_bytes() == (tmp_path / "standard-again/dump.sql").read_bytes()
    assert (tmp_path / "generic/dump.sql").read_bytes() == (tmp_path / "generic-again/dump.sql").read_bytes()
    assert output_rows(tmp_path / "standard")["item"][0]["item"] != output_rows(tmp_path / "generic")["item"][0]["item"]


def test_input_sql_types_preserved_but_model_and_alter_links_ignored(tmp_path):
    source = tmp_path / "data.sql"
    source.write_text("""CREATE TABLE public.sku (sys_id bigint, order_id bigint, note text, qty numeric, day date);
        INSERT INTO public.sku VALUES (777,123,'kg',123.0000001,'2026-10-29');
        ALTER TABLE public.sku ADD PRIMARY KEY (order_id);
        ALTER TABLE public.sku ADD FOREIGN KEY (order_id) REFERENCES public.other(id);
        """, encoding="utf-8")
    output = tmp_path / "out"
    report = run(Options(source, output, tmp_path / "private.snpanon", bypass_model=True))
    row, = output_rows(output)["sku"]
    assert row["order_id"] == "123" and row["qty"] == "123.0000001" and row["day"] == "2026-10-29"
    assert row["note"] != "kg" and row["sys_id"] == ""
    assert not report["relationships_checked"] and not report["keys_checked"]


def test_data_only_sql_without_model_cannot_invent_types(tmp_path):
    source = tmp_path / "data.sql"
    source.write_text("INSERT INTO public.sku(item,loc) VALUES ('SYN_I','SYN_L');", encoding="utf-8")
    with pytest.raises(AnonymizerError):
        run(Options(source, tmp_path / "out", tmp_path / "private.snpanon", bypass_model=True))
    assert not (tmp_path / "out").exists()


def test_skip_bad_inference_keeps_other_tables(tmp_path):
    folder = tmp_path / "source"
    folder.mkdir()
    (folder / "bad.csv").write_text("amount\n42\nSYN_PRIVATE_BAD\n", encoding="utf-8")
    (folder / "good.csv").write_text("note,qty\nSYN_PRIVATE_GOOD,42\n", encoding="utf-8")
    output = tmp_path / "out"
    log = RunLog()
    report = run(Options(folder, output, tmp_path / "private.snpanon", bypass_model=True,
                         skip_error_tables=True, log=log))
    assert set(output_rows(output)) == {"good"}
    assert [t["table"] for t in report["skipped_tables"]] == ["public.bad"]
    assert "SYN_PRIVATE_BAD" not in json.dumps(report)
    assert "SYN_PRIVATE_BAD" not in str(log.entries)


def test_gui_bypass_defaults_off_and_is_available(tk_root):
    from anonymizer.gui import Application
    for child in tk_root.winfo_children():
        child.destroy()
    app = Application(tk_root)
    assert not app.bypass_model.get()
    app.bypass_model.set(True)
    assert app.bypass_model.get()


def test_global_metadata_limit_is_fatal_for_sql_even_with_skip(tmp_path, monkeypatch):
    from anonymizer.generic_input import InputLimitError
    monkeypatch.setattr("anonymizer.store.MAX_INPUT_METADATA", 3000)
    source = tmp_path / "source.sql"
    source.write_text("CREATE TABLE public.a(note text); INSERT INTO public.a VALUES ('SYN_A');"
                      "CREATE TABLE public.b(note text); CREATE TABLE public.c(note text);", encoding="utf-8")
    with pytest.raises(InputLimitError):
        run(Options(source, tmp_path / "out", tmp_path / "private.snpanon",
                    bypass_model=True, skip_error_tables=True))
    assert not (tmp_path / "out").exists()


def test_sql_catalog_is_bounded_before_collecting_all_fields(tmp_path, monkeypatch):
    from anonymizer.generic_input import InputLimitError
    monkeypatch.setattr("anonymizer.sql_reader.MAX_COLUMNS", 2)
    source = tmp_path / "source.sql"
    source.write_text("CREATE TABLE public.a(one text,two text,three text);", encoding="utf-8")
    with pytest.raises(InputLimitError):
        run(Options(source, tmp_path / "out", tmp_path / "private.snpanon", bypass_model=True,
                    skip_error_tables=True))


def test_alter_constraints_do_not_accumulate_without_model(tmp_path):
    from anonymizer.metadata import Model
    from anonymizer.sql_reader import parse_sql
    from anonymizer.store import Store
    source = tmp_path / "source.sql"
    source.write_text("CREATE TABLE public.a(code integer);"
                      + "ALTER TABLE public.a ADD FOREIGN KEY(code) REFERENCES public.b(code);" * 100,
                      encoding="utf-8")
    with Store(tmp_path / "spool.sqlite", Model(), bypass_model=True) as store:
        parse_sql(source, store)
        assert not store.tables[("public", "a")]["fk_groups"]
        assert not store.tables[("public", "a")]["key_groups"]


def test_manifest_limit_is_fatal_before_json_decoding(tmp_path, monkeypatch):
    from anonymizer.generic_input import InputLimitError
    folder = tmp_path / "source"
    folder.mkdir()
    (folder / "custom.csv").write_text("note\nSYN_A\n", encoding="utf-8")
    (folder / "manifest.json").write_text(" " * 50, encoding="utf-8")
    monkeypatch.setattr("anonymizer.csv_reader.MAX_INPUT_MANIFEST", 20)
    with pytest.raises(InputLimitError, match="manifest"):
        run(Options(folder, tmp_path / "out", tmp_path / "private.snpanon", bypass_model=True,
                    skip_error_tables=True))


@pytest.mark.parametrize("true_value,false_value", [("True", "False"), ("tRuE", "fAlSe")])
def test_auto_boolean_case_variants_preserved(tmp_path, true_value, false_value):
    folder = tmp_path / "source"
    folder.mkdir()
    (folder / "custom.csv").write_text(f"flag\n{true_value}\n{false_value}\n", encoding="utf-8")
    output = tmp_path / "out"
    run(Options(folder, output, tmp_path / "private.snpanon", bypass_model=True, delimiter=","))
    assert [row["flag"] for row in output_rows(output)["custom"]] == [true_value, false_value]


def test_unknown_create_kinds_cannot_grow_sql_skip_catalog(tmp_path):
    source = tmp_path / "source.sql"
    source.write_text("CREATE SYN_UNKNOWN_KIND random;", encoding="utf-8")
    with pytest.raises(AnonymizerError, match="CREATE"):
        run(Options(source, tmp_path / "out", tmp_path / "private.snpanon", bypass_model=True,
                    skip_error_tables=True))
