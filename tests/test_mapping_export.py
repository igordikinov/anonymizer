"""Mapping export to a separate Excel file; synthetic data only.

The mapping file carries ORIGINAL values: it must never land inside the
published package (nor overwrite the project/source), and its failure must
block publication (all-or-nothing).
"""
import os
import sqlite3
from threading import Event
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from anonymizer.errors import AnonymizerError, Cancelled
from anonymizer.mapping_export import ROWS_PER_SHEET, cell_value, write_mapping_xlsx
from anonymizer.pipeline import Options, run

SOURCE = """
      CREATE TABLE public.item(item text PRIMARY KEY, note text);
      INSERT INTO public.item VALUES ('SYN_A','SECRET_NOTE'),('SYN_B','OTHER_NOTE');
"""


@pytest.mark.parametrize("error_type", [AnonymizerError, Cancelled, OSError])
def test_error_after_written_row_removes_only_owned_xml_temps(tmp_path, monkeypatch, error_type):
    from pathlib import Path
    from openpyxl.worksheet import _writer

    target = tmp_path / "existing-map.xlsx"
    write_mapping_xlsx(target, [("domain", "SYN_OLD", "Attribute 0001")])
    existing = target.read_bytes()
    foreign = tmp_path / "foreign.xml"
    foreign.write_text("SYN_FOREIGN", encoding="utf-8")
    _writer.ALL_TEMP_FILES.append(str(foreign))
    created = []
    original = _writer.create_temporary_file

    def tracked(*args, **kwargs):
        path = original(*args, **kwargs)
        created.append(path)
        return path

    monkeypatch.setattr(_writer, "create_temporary_file", tracked)

    def rows():
        yield "domain", "SYN_PRIVATE_ORIGINAL", "Attribute 0002"
        raise error_type("synthetic error")

    try:
        with pytest.raises(AnonymizerError):
            write_mapping_xlsx(target, rows())
        assert created
        assert all(not Path(path).exists() and path not in _writer.ALL_TEMP_FILES for path in created)
        assert target.read_bytes() == existing
        assert foreign.read_text(encoding="utf-8") == "SYN_FOREIGN"
        assert str(foreign) in _writer.ALL_TEMP_FILES
    finally:
        _writer.ALL_TEMP_FILES.remove(str(foreign))


def test_mapping_export_contains_all_pairs_outside_package(tmp_path):
    src = tmp_path / "source.sql"
    src.write_text(SOURCE, encoding="utf-8")
    output = tmp_path / "out"
    mapping = tmp_path / "private" / "map.xlsx"
    mapping.parent.mkdir()
    run(Options(src, output, tmp_path / "synthetic.snpanon", mapping_xlsx=mapping))
    assert mapping.is_file() and output.is_dir()
    assert not (output / "map.xlsx").exists() and not list(output.rglob("*.xlsx"))
    book = load_workbook(mapping, read_only=True)
    (sheet,) = book.worksheets
    rows = list(sheet.iter_rows(values_only=True))
    book.close()
    assert rows[0] == ("Домен", "Исходное значение", "Замена")
    pairs = {(row[0], row[1], row[2]) for row in rows[1:]}
    assert any(domain == "schema" and old == "public" and str(new).startswith("snp_a")
               for domain, old, new in pairs)
    assert any(old == "SYN_A" and old != new for _, old, new in pairs)
    assert any(old == "SECRET_NOTE" for _, old, _ in pairs)  # originals live here by design
    # Every project mapping row is exported.
    with sqlite3.connect(tmp_path / "synthetic.snpanon") as db:
        (total,) = db.execute("SELECT count(*) FROM mapping").fetchone()
    assert len(rows) - 1 == total


def test_export_failure_blocks_publication_and_rolls_back(tmp_path):
    src = tmp_path / "source.sql"
    src.write_text(SOURCE, encoding="utf-8")
    project = tmp_path / "synthetic.snpanon"
    target = tmp_path / "sub" / "out.xlsx"  # parent missing -> openpyxl save fails
    with pytest.raises(AnonymizerError, match="карты замен"):
        run(Options(src, tmp_path / "out", project, mapping_xlsx=target))
    assert not (tmp_path / "out").exists()
    with sqlite3.connect(project) as db:
        assert db.execute("SELECT count(*) FROM mapping").fetchone()[0] == 0  # rolled back


def test_mapping_path_equal_to_project_or_output_is_rejected(tmp_path):
    src = tmp_path / "source.sql"
    src.write_text(SOURCE, encoding="utf-8")
    project = tmp_path / "synthetic.snpanon"
    output = tmp_path / "out"
    with pytest.raises(AnonymizerError, match="не может совпадать с проектом"):
        run(Options(src, output, project, mapping_xlsx=project))
    with pytest.raises(AnonymizerError, match="вне папки результата"):
        run(Options(src, output, project, mapping_xlsx=output / "map.xlsx"))
    with pytest.raises(AnonymizerError, match="вне папки результата"):
        run(Options(src, output, project, mapping_xlsx=output))
    assert not output.exists()
    with sqlite3.connect(project) as db:
        assert db.execute("PRAGMA integrity_check").fetchone()[0] == "ok"


def test_existing_mapping_file_is_not_overwritten_without_permission(tmp_path):
    src = tmp_path / "source.sql"
    src.write_text(SOURCE, encoding="utf-8")
    mapping = tmp_path / "map.xlsx"
    mapping.write_bytes(b"previous user data")
    with pytest.raises(AnonymizerError, match="уже существует"):
        run(Options(src, tmp_path / "out", tmp_path / "synthetic.snpanon", mapping_xlsx=mapping))
    assert mapping.read_bytes() == b"previous user data"
    assert not (tmp_path / "out").exists()
    run(Options(src, tmp_path / "out", tmp_path / "synthetic.snpanon",
                mapping_xlsx=mapping, mapping_overwrite=True))
    assert mapping.read_bytes().startswith(b"PK")  # a real xlsx replaced it


def test_cancel_during_export_propagates_and_leaves_no_output(tmp_path):
    target = tmp_path / "map.xlsx"

    def rows():
        for i in range(10):
            if i == 3:
                raise AnonymizerError("Обработка отменена. Итоговый пакет не создан.")
            yield (f"d{i}", f"o{i}", f"n{i}")

    with pytest.raises(AnonymizerError, match="отменена"):
        write_mapping_xlsx(target, rows())
    assert not target.exists()  # failed save leaves no partial file


def test_no_flag_means_no_export(tmp_path):
    src = tmp_path / "source.sql"
    src.write_text(SOURCE, encoding="utf-8")
    output = tmp_path / "out"
    run(Options(src, output, tmp_path / "synthetic.snpanon"))
    assert output.is_dir() and not list(tmp_path.rglob("*.xlsx"))


def test_cell_value_truncates_and_cleans_control_characters():
    assert cell_value("ok") == "ok"
    assert "\x00" not in cell_value("bad\x00\x01value")
    long = "x" * 40_000
    result = cell_value(long)
    assert len(result) <= 32_000 and result.endswith("...[truncated]")


@pytest.mark.parametrize("value", ["=1+1", "#N/A", "+123", "@SUM(A1)", "-123", "00123"])
def test_mapping_cells_preserve_values_as_text(tmp_path, value):
    target = tmp_path / "text.xlsx"
    write_mapping_xlsx(target, [(value, value, value)])
    book = load_workbook(target, read_only=True, data_only=False)
    try:
        (header, row) = list(book.active.iter_rows())
        assert all(cell.data_type == "s" for cell in (*header, *row))
        assert [cell.value for cell in row] == [value, value, value]
    finally:
        book.close()


def test_sheet_rollover_on_overflow(tmp_path):
    rows = ((f"domain{i}", f"old{i}", f"new{i}") for i in range(25))
    target = tmp_path / "overflow.xlsx"
    with patch.object(__import__("anonymizer.mapping_export", fromlist=["x"]), "ROWS_PER_SHEET", 10):
        assert write_mapping_xlsx(target, rows) == 25
    book = load_workbook(target, read_only=True)
    assert [sheet.title for sheet in book.worksheets] == ["mapping", "mapping_2", "mapping_3"]
    counts = [sum(1 for _ in sheet.iter_rows(values_only=True)) - 1 for sheet in book.worksheets]
    book.close()
    assert counts == [10, 10, 5]
