"""Synthetic Latin-generation acceptance, including unchanged private history."""
import csv
from datetime import date
import json
import sqlite3

import pytest
from openpyxl import Workbook, load_workbook

from anonymizer.errors import AnonymizerError
from anonymizer.mapping import INDUSTRIES, Project
from anonymizer.pipeline import Options, run


def exported_rows(output):
    manifest = json.loads((output / "csv/manifest.json").read_text(encoding="utf-8"))
    result = {}
    for table in manifest["tables"]:
        with (output / "csv" / table["file"]).open(encoding="utf-8", newline="") as stream:
            result[table["table"]] = list(csv.DictReader(stream))
    return result


def seed_mapping(path, domain, old, new):
    project = Project(path, "food")
    project.db.execute("INSERT INTO mapping VALUES (?,?,?,?)", (domain, old, new, 1))
    project.finish(True)


@pytest.mark.parametrize("industry", tuple(INDUSTRIES))
def test_all_industry_products_and_description_kinds_generate_ascii(tmp_path, industry):
    project = Project(tmp_path / "private.snpanon", industry)
    try:
        # Include the wraparound item, so both the pool and ordinal suffix matter.
        products = [project.pseudonym("descriptions:item", f"Синтетика {n}", label="item")
                    for n in range(1, len(INDUSTRIES[industry]["products"]) + 2)]
        assert len(set(products)) == len(products)
        assert all(value.isascii() and "- series " in value for value in products)
        for kind in ("loc", "res", "prodsource_sp", "demand_stream", "brand", "category",
                     "supplier", "planner", "address", "text", "unknown_kind"):
            value = project.pseudonym("descriptions:" + kind, "Синтетический текст", label=kind)
            assert value.isascii() and any(char.isalpha() for char in value)
        assert all(new.isascii() for _, _, new in project.iter_mapping())
    finally:
        project.finish(False)


def test_standard_fk_descriptions_and_json_are_latin_but_preserved_values_stay(tmp_path):
    source = tmp_path / "source.sql"
    payload = json.dumps({"description": "Синтетическое описание", "Свойство": "Значение",
                          "qty": 42.125, "date": "2026-10-29", "enabled": True}, ensure_ascii=False)
    source.write_text(
        "CREATE TABLE public.item(sys_id bigint, item text PRIMARY KEY, proddescr text, "
        "baseuom text, payload jsonb, quantity numeric, day date, active boolean, "
        "update_date_time text, change_author text);"
        "CREATE TABLE public.loc(loc text PRIMARY KEY, locdescr text);"
        "CREATE TABLE public.sku(item text REFERENCES public.item(item), "
        "loc text REFERENCES public.loc(loc), planner text);"
        "INSERT INTO public.item VALUES ('Скрытый ID','КОД','Описание','ШТ','" + payload +
        "',1234567890.123456789,'2026-10-29',true,'Сырое время','Автор');"
        "INSERT INTO public.loc VALUES ('МЕСТО','Площадка');"
        "INSERT INTO public.sku VALUES ('КОД','МЕСТО','Планировщик');", encoding="utf-8")
    output = tmp_path / "out"
    project = tmp_path / "private.snpanon"
    report = run(Options(source, output, project))
    tables = exported_rows(output)
    item, loc, sku = tables["item"][0], tables["loc"][0], tables["sku"][0]
    assert sku["item"] == item["item"] and sku["loc"] == loc["loc"]
    assert all(value.isascii() for value in (item["item"], item["proddescr"], loc["loc"],
                                            loc["locdescr"], sku["planner"]))
    assert item["baseuom"] == "ШТ" and item["update_date_time"] == "Сырое время"
    assert item["sys_id"] == "" and item["change_author"] == "snp"
    assert item["quantity"] == "1234567890.123456789" and item["day"] == "2026-10-29"
    assert item["active"] == "true" and report["validation"]["preserved_fields_equal"]
    changed_json = json.loads(item["payload"])
    assert all(key.isascii() for key in changed_json)
    assert all(not isinstance(value, str) or value.isascii() for value in changed_json.values())
    assert changed_json["qty"] == 42.125 and changed_json["enabled"] is True
    with sqlite3.connect(project) as db:
        mappings = db.execute("SELECT old,new FROM mapping").fetchall()
    assert any(not old.isascii() for old, _ in mappings)
    assert all(new.isascii() for _, new in mappings)


@pytest.mark.parametrize("kind", ["csv", "xlsx"])
def test_generic_unicode_sources_keep_column_prefixes_and_system_policy(tmp_path, kind):
    folder = tmp_path / "source"
    folder.mkdir()
    columns = ["product", "resource", "location", "bomnum", "quantity", "day", "active",
               "sys_id", "update_date_time", "change_author"]
    values = ["Синтетика"] * 4 + [42.125, date(2026, 10, 29), True,
                                 "Исходный ID", "Сырое время", None]
    if kind == "csv":
        with (folder / "optimizer_production.csv").open("w", encoding="utf-8", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(columns)
            writer.writerow(values[:6] + ["true", *values[7:]])
    else:
        book = Workbook(iso_dates=True)
        book.active.append(columns)
        book.active.append(values)
        book.save(folder / "optimizer_production.xlsx")
        book.close()
    output = tmp_path / "out"
    report = run(Options(folder, output, tmp_path / "private.snpanon", bypass_model=True))
    row, = exported_rows(output)["optimizer_production"]
    assert [row[column] for column in columns[:4]] == [f"{column}_000001" for column in columns[:4]]
    assert row["quantity"] == "42.125" and row["day"] == "2026-10-29" and row["active"] == "true"
    assert row["sys_id"] == "" and row["update_date_time"] == "Сырое время"
    assert row["change_author"] == "snp" and not report["relationships_checked"]


def test_ascii_legacy_mapping_is_reused_and_new_values_repeat(tmp_path):
    project = tmp_path / "private.snpanon"
    seed_mapping(project, "generic:public:custom.note", "Старый", "OLD_LATIN_LABEL")
    folder = tmp_path / "source"
    folder.mkdir()
    (folder / "custom.csv").write_text("note,quantity\nСтарый,1\nНовый,2\n", encoding="utf-8")
    dumps = []
    for index in range(2):
        output = tmp_path / f"out-{index}"
        run(Options(folder, output, project, bypass_model=True))
        first, second = exported_rows(output)["custom"]
        assert first["note"] == "OLD_LATIN_LABEL" and second["note"] == "note_000002"
        dumps.append((output / "dump.sql").read_bytes())
    assert dumps[0] == dumps[1]


@pytest.mark.parametrize("bypass", [False, True])
def test_nonascii_legacy_use_fails_without_publication_or_rewriting(tmp_path, bypass):
    project = tmp_path / "private.snpanon"
    domain = ("generic:" if bypass else "") + "public:custom.note"
    seed_mapping(project, domain, "Исходник", "Старая замена")
    source = tmp_path / "source.sql"
    source.write_text("CREATE TABLE public.custom(note text); INSERT INTO public.custom VALUES ('Исходник');",
                      encoding="utf-8")
    output = tmp_path / "out"
    with pytest.raises(AnonymizerError) as caught:
        run(Options(source, output, project, bypass_model=bypass))
    assert "проект" in str(caught.value).lower()
    assert "Исходник" not in str(caught.value) and "Старая замена" not in str(caught.value)
    assert not output.exists()
    with sqlite3.connect(project) as db:
        assert db.execute("SELECT domain,old,new FROM mapping").fetchall() == [(domain, "Исходник", "Старая замена")]


def test_cache_and_new_candidate_cannot_bypass_latin_guard(tmp_path):
    project = Project(tmp_path / "private.snpanon", "food")
    key = ("generic:public:custom.note", "Исходник")
    project.cache[key] = ("Старый кеш", 1)
    try:
        with pytest.raises(AnonymizerError):
            project.pseudonym(*key, avoid_original=True, text_prefix="note")
        with pytest.raises(AnonymizerError):
            project.pseudonym("generic:public:custom.other", "Другой", avoid_original=True,
                              text_prefix="Кириллица")
        assert not project.db.execute("SELECT new FROM mapping").fetchall()
        assert ("generic:public:custom.other", "Другой") not in project.cache
    finally:
        project.finish(False)


def test_unused_nonlatin_history_blocks_mapping_export_but_is_not_rewritten(tmp_path):
    project = tmp_path / "private.snpanon"
    seed_mapping(project, "unused:history", "История", "Старая замена")
    folder = tmp_path / "source"
    folder.mkdir()
    (folder / "custom.csv").write_text("note,quantity\nНовый,1\n", encoding="utf-8")
    output, mapping = tmp_path / "out", tmp_path / "private-map.xlsx"
    with pytest.raises(AnonymizerError):
        run(Options(folder, output, project, bypass_model=True, mapping_xlsx=mapping))
    assert not output.exists() and not mapping.exists()
    with sqlite3.connect(project) as db:
        assert db.execute("SELECT domain,old,new FROM mapping").fetchall() == [
            ("unused:history", "История", "Старая замена")]
    # Unused history need not be scanned when it is not being exported.
    clean_output = tmp_path / "clean-out"
    run(Options(folder, clean_output, project, bypass_model=True))
    assert exported_rows(clean_output)["custom"][0]["note"] == "note_000001"


def test_mapping_export_truncation_keeps_replacement_latin_and_original_unicode(tmp_path):
    project = tmp_path / "private.snpanon"
    original, replacement = "Синтетический исходник", "A" * 33_000
    seed_mapping(project, "unused:history", original, replacement)
    folder = tmp_path / "source"
    folder.mkdir()
    (folder / "custom.csv").write_text("note,quantity\nНовый,1\n", encoding="utf-8")
    mapping = tmp_path / "private-map.xlsx"
    run(Options(folder, tmp_path / "out", project, bypass_model=True, mapping_xlsx=mapping))
    book = load_workbook(mapping, read_only=True)
    try:
        row, = [row for row in book.active.iter_rows(min_row=2, values_only=True)
                if row[0] == "unused:history"]
        assert row[1] == original
        assert row[2].isascii() and len(row[2]) < len(replacement)
        assert "truncated" in row[2]
    finally:
        book.close()
