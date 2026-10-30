"""Readable private map labels; project identities and pairs stay unchanged."""
from contextlib import contextmanager
import json
import sqlite3
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from anonymizer.connectivity import compile_contract, load_profile
from anonymizer.errors import AnonymizerError, Cancelled
from anonymizer.mapping import Project
from anonymizer.mapping_export import HEADER, LABEL_BYTES, LABEL_MARKER, iter_mapping_rows, write_mapping_xlsx
from anonymizer.metadata import Model
from anonymizer.pipeline import Options, run
from anonymizer.store import Store


LEGACY_RESOURCE_KEY = ("connectivity:1d82586f2ce01a65c5f1649152fdd407f30e162dccb8ec6012bf4f32eb070867:"
                "04630553c1f17f8e606bc0ce4eb827cf1759308d65a392c5502d4ed7eca8c71f")
RESOURCE_TABLES = ("optimizer_production", "optimizer_storage", "optimizer_transportation")


def custom_profile(*, shared=False):
    return {"format": "snp-anonymizer-connectivity-v1", "profile_id": "synthetic-assets", "version": "1",
            "namespace_rule": "shared" if shared else "source_schema",
            "domains": [{"id": "RESOURCE", "kind": "text", "prefix": "ASSET",
                         "members": [{"table": "assets", "column": "code"}]}],
            "relations": [], "keys": [], "required_checks": ["domain:RESOURCE"],
            "missing_target_policy": "not_checked"}


@contextmanager
def project_for(tmp_path, project_path, tables, profile="builtin:optimizer-v2"):
    with Store(tmp_path / f"input-{uuid4().hex}.sqlite", Model(), bypass_model=True) as store:
        for schema, table, column, values in tables:
            name = (schema, table)
            store.define(name, [{"name": column, "sql_type": "text"}])
            for value in values:
                store.add_row(name, [column], [value])
        contract = compile_contract(store, load_profile(profile))
        project = Project(project_path, "food", contract=contract)
        try:
            project.prepare_domains(store)
            yield project
        except BaseException:
            project.finish(False)
            raise
        else:
            project.finish(True)


def resources(schema="public", tables=RESOURCE_TABLES, values=("SYN_A", "SYN_B")):
    return [(schema, table, "resource", values) for table in tables]


def test_split_resource_domains_are_readable_without_changing_project(tmp_path):
    with project_for(tmp_path, tmp_path / "private.snpanon", resources()) as project:
        raw = list(project.iter_mapping())
        before_settings = project.db.execute("SELECT * FROM settings ORDER BY name").fetchall()
        before_bindings = project.db.execute("SELECT * FROM connectivity_bindings ORDER BY node").fetchall()
        shown = list(iter_mapping_rows(project))
        keys = {row[0] for row in raw}
        assert len(keys) == 3 and LEGACY_RESOURCE_KEY not in keys
        assert {(old, new) for _, old, new in shown} == {(old, new) for _, old, new in raw}
        assert len(shown) == len(raw) == 6
        labels = {row[0] for row in shown}
        assert len(labels) == 3
        for identifier, table in (("RESOURCE_PROD", "optimizer_production"),
                                  ("RESOURCE_STOCK", "optimizer_storage"),
                                  ("RESOURCE_TRANSP", "optimizer_transportation")):
            label, = [label for label in labels if "(" + identifier + ")" in label]
            assert "схема public:" in label and table + ".resource" in label
            assert all(other + ".resource" not in label for other in RESOURCE_TABLES if other != table)
            assert "connectivity:" not in label and "1d82586f" not in label
        assert list(project.iter_mapping()) == raw
        assert project.db.execute("SELECT * FROM settings ORDER BY name").fetchall() == before_settings
        assert project.db.execute("SELECT * FROM connectivity_bindings ORDER BY node").fetchall() == before_bindings
        assert dict(before_settings)["format"] == "snp-anonymizer-project-v2"


def test_pipeline_excel_keeps_three_columns_text_cells_and_pairs_with_readable_labels(tmp_path):
    source = tmp_path / "source.sql"
    source.write_text("\n".join(
        f"CREATE TABLE public.{table} (resource text);\n"
        f"INSERT INTO public.{table} VALUES ('=1+1'),('SYN_B');"
        for table in RESOURCE_TABLES), encoding="utf-8")
    path, workbook = tmp_path / "private.snpanon", tmp_path / "private-map.xlsx"
    result = run(Options(source, tmp_path / "out", path, bypass_model=True,
                         connectivity_profile="builtin:optimizer-v2", mapping_xlsx=workbook))
    assert result["dataset_published"]
    assert not list((tmp_path / "out").rglob("*.xlsx"))
    with sqlite3.connect(path) as db:
        raw = db.execute("SELECT domain,old,new FROM mapping ORDER BY domain,ordinal").fetchall()
    book = load_workbook(workbook)
    try:
        sheet, = book.worksheets
        data = list(sheet.iter_rows())
        assert tuple(c.value for c in data[0]) == HEADER
        assert len(data) - 1 == len(raw)
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref == f"A1:C{len(data)}"
        assert sheet.column_dimensions["A"].width >= 60
        for row in data[1:]:
            assert all(cell.data_type == "s" for cell in row)
            assert all(cell.alignment.wrap_text for cell in row)
        shown = [tuple(c.value for c in row) for row in data[1:]]
        assert any("(RESOURCE_PROD)" in label and old == "=1+1" for label, old, _ in shown)
        assert {(old, new) for label, old, new in shown if label.startswith("Ресурс")}
        assert {(old, new) for label, old, new in shown if label.startswith("Ресурс")} == {
            (old, new) for domain, old, new in raw if domain.startswith("connectivity:")}
    finally:
        book.close()


def test_reopened_subset_keeps_historical_schemas_and_fields(tmp_path):
    path = tmp_path / "private.snpanon"
    with project_for(tmp_path, path, resources("alpha") + resources("beta")) as project:
        raw = list(project.iter_mapping())
        original_labels = {label for label, _, _ in iter_mapping_rows(project)}
    with project_for(tmp_path, path, resources("beta", ("optimizer_storage",))) as project:
        shown = list(iter_mapping_rows(project))
        assert list(project.iter_mapping()) == raw
        assert {label for label, _, _ in shown} == original_labels
        assert any("схема alpha" in label for label, _, _ in shown)
        assert len({label for label, _, _ in shown}) == 6
        assert all(sum(f"{table}.resource" in label for table in RESOURCE_TABLES) == 1 for label, _, _ in shown)


def test_shared_namespace_lists_full_paths_and_custom_profile_does_not_borrow_builtin_meaning(tmp_path):
    path = tmp_path / "private.snpanon"
    tables = [(schema, "assets", "code", ("SYN_A",)) for schema in ("alpha", "beta")]
    with project_for(tmp_path, path, tables, custom_profile(shared=True)) as project:
        raw = list(project.iter_mapping())
        shown = list(iter_mapping_rows(project))
        assert len(raw) == len(shown) == 1
        label = shown[0][0]
        assert label.startswith("Домен RESOURCE, общий namespace:")
        assert "alpha.assets.code" in label and "beta.assets.code" in label
        assert "Ресурс" not in label
    with project_for(tmp_path, path, tables[:1], custom_profile(shared=True)) as project:
        assert list(iter_mapping_rows(project)) == shown


def test_external_source_schema_namespace_stays_distinguishable(tmp_path):
    tables = [(schema, "assets", "code", ("SYN_A",)) for schema in ("alpha", "beta")]
    with project_for(tmp_path, tmp_path / "private.snpanon", tables, custom_profile()) as project:
        labels = {label for label, _, _ in iter_mapping_rows(project)}
        assert len(labels) == 2
        assert any("схема alpha" in label for label in labels)
        assert any("схема beta" in label for label in labels)
        assert all(label.startswith("Домен RESOURCE") for label in labels)


def test_legacy_domain_labels_and_raw_iterator_are_unchanged(tmp_path):
    project = Project(tmp_path / "private.snpanon", "food")
    try:
        project.pseudonym("generic:public.assets.code", "SYN_A")
        project.pseudonym("public:item.item", "SYN_B")
        raw = list(project.iter_mapping())
        assert list(iter_mapping_rows(project)) == raw
    finally:
        project.finish(False)


def test_readable_source_namespace_and_originals_stay_only_in_private_map(tmp_path):
    source = tmp_path / "source.sql"
    source.write_text("CREATE TABLE private_client_ns.assets (code text);\n"
                      "INSERT INTO private_client_ns.assets VALUES ('SYN_PRIVATE_ID');", encoding="utf-8")
    workbook, output = tmp_path / "private-map.xlsx", tmp_path / "out"
    result = run(Options(source, output, tmp_path / "private.snpanon", bypass_model=True,
                         connectivity_profile=custom_profile(), require_integrity=True, mapping_xlsx=workbook))
    assert result["integrity_status"] == "PASSED"
    book = load_workbook(workbook, read_only=True)
    try:
        shown = list(book.active.iter_rows(values_only=True))[1:]
        assert any("схема private_client_ns" in label and old == "SYN_PRIVATE_ID" for label, old, _ in shown)
    finally:
        book.close()
    for path in output.rglob("*"):
        if path.is_file():
            body = path.read_bytes()
            assert b"private_client_ns" not in body and b"SYN_PRIVATE_ID" not in body


@pytest.mark.parametrize("damage", ["invalid_node", "invalid_semantic", "null_node", "wrong_namespace", "missing_bindings"])
def test_corrupted_bindings_fail_with_safe_error(tmp_path, damage):
    with project_for(tmp_path, tmp_path / "private.snpanon", resources(values=("SYN_PRIVATE",))) as project:
        if damage == "invalid_node":
            project.db.execute("UPDATE connectivity_bindings SET node='SYN_PRIVATE_BAD' WHERE node=(SELECT min(node) FROM connectivity_bindings)")
        elif damage == "null_node":
            project.db.execute("UPDATE connectivity_bindings SET node=NULL WHERE node=(SELECT min(node) FROM connectivity_bindings)")
        elif damage == "invalid_semantic":
            project.db.execute("UPDATE connectivity_bindings SET semantic='SYN_PRIVATE_BAD'")
        elif damage == "wrong_namespace":
            semantic = json.loads(project.db.execute("SELECT semantic FROM connectivity_bindings LIMIT 1").fetchone()[0])
            semantic["namespace"] = "private_wrong"
            project.db.execute("UPDATE connectivity_bindings SET semantic=?", (json.dumps(semantic),))
        else:
            project.db.execute("DELETE FROM connectivity_bindings")
        with pytest.raises(AnonymizerError, match="привязки") as error:
            list(iter_mapping_rows(project))
        assert "SYN_PRIVATE" not in str(error.value) and "private_wrong" not in str(error.value)


def add_historical_bindings(project, count):
    domain, semantic = project.db.execute("SELECT domain,semantic FROM connectivity_bindings LIMIT 1").fetchone()
    project.db.executemany("INSERT INTO connectivity_bindings VALUES (?,?,?)", (
        (json.dumps(["public", f"historical_assets_{index:06d}", "resource"]), domain, semantic)
        for index in range(count)))


def test_long_history_is_bounded_explicitly_shortened_and_still_fully_consumed(tmp_path):
    with project_for(tmp_path, tmp_path / "private.snpanon", resources(values=("SYN_A",))) as project:
        add_historical_bindings(project, 4000)
        count = 0

        def check():
            nonlocal count
            count += 1

        rows = list(iter_mapping_rows(project, check))
        label, = {label for label, _, _ in rows if label.endswith(LABEL_MARKER)}
        assert all(len(label.encode("utf-8")) <= LABEL_BYTES for label, _, _ in rows)
        assert label.endswith(LABEL_MARKER)
        assert count > 4000  # truncation does not stop the cancellable stream


def test_cancel_during_registry_sort_propagates_and_releases_sqlite_cursors(tmp_path):
    with project_for(tmp_path, tmp_path / "private.snpanon", resources(values=("SYN_A",))) as project:
        add_historical_bindings(project, 4000)
        calls = 0
        statements = []
        project.db.set_trace_callback(statements.append)

        def cancel():
            nonlocal calls
            calls += 1
            raise Cancelled("synthetic cancellation")

        with pytest.raises(Cancelled):
            list(iter_mapping_rows(project, cancel))
        assert calls >= 1
        assert any("FROM connectivity_bindings ORDER BY" in statement for statement in statements)
        assert not any("FROM mapping ORDER BY" in statement for statement in statements)
        # A retained cursor would lock this table; a retained handler would
        # interrupt the long statement even after cancellation was handled.
        project.db.execute("DROP TABLE connectivity_bindings")
        assert project.db.execute("WITH RECURSIVE n(x) AS (VALUES(0) UNION ALL SELECT x+1 FROM n WHERE x<5000) SELECT sum(x) FROM n").fetchone()[0] == 12502500


def test_explicit_generator_close_releases_registry_cursor(tmp_path):
    with project_for(tmp_path, tmp_path / "private.snpanon", resources()) as project:
        rows = iter_mapping_rows(project)
        next(rows)
        rows.close()
        project.db.execute("DROP TABLE connectivity_bindings")


def test_filter_and_freeze_are_correct_across_rollover_and_empty_book(tmp_path, monkeypatch):
    import anonymizer.mapping_export as module
    monkeypatch.setattr(module, "ROWS_PER_SHEET", 2)
    target = tmp_path / "rolled.xlsx"
    write_mapping_xlsx(target, (("domain", str(index), "A") for index in range(5)))
    book = load_workbook(target)
    try:
        assert [sheet.auto_filter.ref for sheet in book] == ["A1:C3", "A1:C3", "A1:C2"]
        assert all(sheet.freeze_panes == "A2" for sheet in book)
    finally:
        book.close()
    empty = tmp_path / "empty.xlsx"
    write_mapping_xlsx(empty, [])
    book = load_workbook(empty)
    try:
        sheet, = book.worksheets
        assert sheet.auto_filter.ref == "A1:C1" and sheet.freeze_panes == "A2"
    finally:
        book.close()
