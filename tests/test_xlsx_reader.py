"""Synthetic OOXML fixtures, including exact lexemes Excel writers would round."""
from pathlib import Path
import sqlite3
import struct
import zipfile
from xml.sax.saxutils import escape

import pytest
from openpyxl import Workbook

from anonymizer.errors import AnonymizerError, Cancelled
from anonymizer.metadata import Model
from anonymizer.store import Store
from anonymizer.xlsx_reader import parse_xlsx, XLSXLimitError, _date_value, _Strings
import anonymizer.xlsx_reader as reader

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG = "http://schemas.openxmlformats.org/package/2006/relationships"


def model():
    result = Model()
    result.tables["contract"] = {"fields": [
        {"name": name, "type": kind, "nullable": nullable, "size": None, "keys": [], "fk": None}
        for name, kind, nullable in [("code", "TEXT", False), ("note", "TEXT", True),
            ("qty", "NUMBER", True), ("integer", "INTEGER", True), ("day", "DATE", True),
            ("stamp", "DATE_TIME", True), ("clock", "TIME", True), ("flag", "BOOLEAN", True)]]}
    return result


def cell(ref, value=None, kind="inlineStr"):
    if value is None:
        return f'<c r="{ref}"/>'
    content = f'<is><t>{escape(value)}</t></is>' if kind == "inlineStr" else f'<v>{escape(value)}</v>'
    return f'<c r="{ref}" t="{kind}">{content}</c>'


def book(path, sheet, *, shared=None, epoch=False, sheets=1, extra=None):
    parts = {
        "_rels/.rels": f'<Relationships xmlns="{PKG}"><Relationship Id="root" Type="{REL}/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        "xl/workbook.xml": f'<workbook xmlns="{NS}" xmlns:r="{REL}"><workbookPr date1904="{int(epoch)}"/><sheets>' + ''.join(f'<sheet name="Arbitrary{i}" sheetId="{i}" r:id="sheet{i}"/>' for i in range(1,sheets+1)) + '</sheets></workbook>',
        "xl/_rels/workbook.xml.rels": f'<Relationships xmlns="{PKG}"><Relationship Id="sheet1" Type="{REL}/worksheet" Target="worksheets/sheet1.xml"/>' + (f'<Relationship Id="strings" Type="{REL}/sharedStrings" Target="sharedStrings.xml"/>' if shared is not None else '') + '</Relationships>',
        "xl/worksheets/sheet1.xml": sheet,
    }
    if shared is not None:
        parts["xl/sharedStrings.xml"] = f'<sst xmlns="{NS}">' + ''.join(f'<si><t>{escape(s)}</t></si>' for s in shared) + '</sst>'
    parts.update(extra or {})
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, value in parts.items():
            archive.writestr(name, value)
    return path


def sheet(rows, tail=""):
    return f'<worksheet xmlns="{NS}"><sheetData>{rows}</sheetData>{tail}</worksheet>'


def row(number, *cells, attrs=""):
    return f'<row r="{number}" {attrs}>' + ''.join(cells) + '</row>'


def read(tmp_path, source, *, skip=False, check=lambda: None):
    with Store(tmp_path / ('skip.sqlite' if skip else 'store.sqlite'), model()) as store:
        parse_xlsx(source, store, check=check, skip_error_tables=skip)
        if not store.tables:
            return [], store.count, store.dropped_tables
        entry = store.tables['public', 'contract']
        rows = [dict(zip(store.columns(('public','contract')), values)) for values in store.rows(entry['id'])]
        return rows, store.count, store.dropped_tables


def test_exact_numbers_sparse_hidden_rows_and_nulls(tmp_path):
    header=row(1,*[cell(f'{chr(65+i)}1',name) for i,name in enumerate(['code','note','qty','integer','flag'])])
    data=row(2,cell('A2','000123'),cell('B2',''),cell('C2','12345678901234567890.123456789','n'),cell('E2','1','b'),attrs='hidden="1"')
    data+=row(3,cell('A3','SYN_B'),cell('B3',r'\N'),cell('C3',''),cell('D3','0','n'))
    data+=row(4,cell('A4'),cell('E4'))
    source=book(tmp_path/'contract_20261028140000.xlsx',sheet(header+data,'<autoFilter ref="A1:E4"/>'))
    rows,count,_=read(tmp_path,source)
    assert count==2
    assert rows[0]['code']=='000123' and rows[0]['note']==''
    assert rows[0]['qty']=='12345678901234567890.123456789' and rows[0]['integer'] is None
    assert rows[0]['flag']=='true'
    assert rows[1]['qty'] is None and rows[1]['note'] is None and rows[1]['integer']=='0'


def test_regular_openpyxl_workbook_is_supported(tmp_path):
    source=tmp_path/'contract.xlsx'
    workbook=Workbook()
    workbook.active.title='Произвольное имя'
    workbook.active.append(['code','note','qty'])
    workbook.active.append(['SYN_CODE','  Юникод  ',0])
    workbook.save(source)
    rows,count,_=read(tmp_path,source)
    assert count==1 and rows[0]['note']=='  Юникод  ' and rows[0]['qty']=='0'


def test_shared_strings_and_rich_text(tmp_path):
    header=row(1,cell('A1','0','s'),cell('B1','1','s'))
    data=row(2,cell('A2','2','s'),cell('B2','3','s'))
    sst=f'<sst xmlns="{NS}"><si><t>code</t></si><si><t>note</t></si><si><t>SYN_CODE</t></si><si><r><t> rich </t></r><r><t>text</t></r><rPh sb="0" eb="1"><t>PHONETIC</t></rPh></si></sst>'
    source=book(tmp_path/'contract.xlsx',sheet(header+data),shared=[],extra={'xl/sharedStrings.xml':sst})
    rows,_,_=read(tmp_path,source)
    assert rows[0]['note']==' rich text'


@pytest.mark.parametrize('epoch,serial,expected',[(False,'0','1899-12-30'),(False,'1','1900-01-01'),(False,'59','1900-02-28'),(False,'61','1900-03-01'),(True,'0','1904-01-01')])
def test_numeric_dates_obey_epoch(tmp_path,epoch,serial,expected):
    header=row(1,cell('A1','code'),cell('B1','day'),cell('C1','stamp'),cell('D1','clock'))
    data=row(2,cell('A2','SYN'),cell('B2',serial,'n'),cell('C2',str(int(serial))+'.5','n'),cell('D2','0.5','n'))
    source=book(tmp_path/'contract.xlsx',sheet(header+data),epoch=epoch)
    rows,_,_=read(tmp_path,source)
    assert rows[0]['day']==expected
    assert rows[0]['stamp']==expected+'T12:00:00'
    assert rows[0]['clock']=='12:00:00'


@pytest.mark.parametrize('serial,kind',[('60','DATE'),('60.5','DATE_TIME'),('0.0000000000000000001','TIME'),('1.1','DATE'),('NaN','DATE'),('1e999999','DATE'),('1e-999999999','TIME')])
def test_bad_dates_never_round(serial,kind):
    with pytest.raises(AnonymizerError):
        _date_value(serial,kind,False)


def test_iso_dates_and_explicit_null(tmp_path):
    header=row(1,cell('A1','code'),cell('B1','day'),cell('C1','stamp'),cell('D1','clock'))
    data=row(2,cell('A2','SYN'),cell('B2','2026-10-28','d'),cell('C2','2026-10-28T13:14:15.000001','d'),cell('D2','01:02:03.000004'))
    rows,_,_=read(tmp_path,book(tmp_path/'contract.xlsx',sheet(header+data)))
    assert rows[0]['stamp']=='2026-10-28T13:14:15.000001' and rows[0]['clock']=='01:02:03.000004'


@pytest.mark.parametrize('bad',[
    '<c r="B3"><f>PRIVATE_FORMULA</f><v>10</v></c>',
    '<c r="B3"><f>PRIVATE_FORMULA</f></c>',
    '<c r="B3" t="e"><v>PRIVATE_ERROR</v></c>',
])
def test_formulas_and_errors_in_unknown_columns_are_safe_and_skip_is_atomic(tmp_path,bad):
    header=row(1,cell('A1','code'),cell('B1','unknown'))
    data=row(2,cell('A2','SYN_ONE'))+row(3,cell('A3','SYN_TWO'),bad)
    source=book(tmp_path/'contract.xlsx',sheet(header+data))
    with pytest.raises(AnonymizerError) as error:
        read(tmp_path,source)
    assert 'B3' in str(error.value) and 'contract' in str(error.value)
    assert 'PRIVATE' not in str(error.value)
    rows,count,dropped=read(tmp_path,source,skip=True)
    assert rows==[] and count==0 and dropped[0][0]=='public.contract'


@pytest.mark.parametrize('tail,sheets', [('<mergeCells><mergeCell ref="A2:B2"/></mergeCells>',1),('',2)])
def test_late_merge_and_extra_sheet_reject_or_drop_table(tmp_path,tail,sheets):
    source=book(tmp_path/'contract.xlsx',sheet(row(1,cell('A1','code'))+row(2,cell('A2','SYN')),tail),sheets=sheets)
    rows,count,dropped=read(tmp_path,source,skip=True)
    assert count==0 and dropped


@pytest.mark.parametrize('header', [row(1,cell('A1','code'),cell('B1','code')),row(2,cell('A2','code')),row(1,cell('A1','note'))])
def test_invalid_header_or_missing_required_field(tmp_path,header):
    source=book(tmp_path/'contract.xlsx',sheet(header+row(3,cell('A3','SYN'))))
    with pytest.raises(AnonymizerError):
        read(tmp_path,source)


@pytest.mark.parametrize('malicious',[
    '<!DOCTYPE worksheet [<!ENTITY x "PRIVATE">]>',
    '<!DOCTYPE worksheet SYSTEM "file:///PRIVATE">',
])
def test_dtd_and_entities_are_rejected(tmp_path,malicious):
    source=book(tmp_path/'contract.xlsx',malicious+sheet(row(1,cell('A1','code'))))
    with pytest.raises(AnonymizerError,match='DTD'):
        read(tmp_path,source)


def test_external_relationships_are_never_resolved(tmp_path):
    external=f'<Relationships xmlns="{PKG}"><Relationship Id="bad" Type="{REL}/hyperlink" Target="https://PRIVATE.invalid/" TargetMode="External"/></Relationships>'
    source=book(tmp_path/'contract.xlsx',sheet(row(1,cell('A1','code'))),extra={'xl/worksheets/_rels/sheet1.xml.rels':external})
    with pytest.raises(AnonymizerError,match='внешние ссылки'):
        read(tmp_path,source)


@pytest.mark.parametrize('setting', ['MAX_UNCOMPRESSED','MAX_DIRECTORY','MAX_METADATA','MAX_FIELD','MAX_ENTRIES'])
def test_resource_violations_are_fatal_even_in_skip_mode(tmp_path,monkeypatch,setting):
    source=book(tmp_path/'contract.xlsx',sheet(row(1,cell('A1','code'))+row(2,cell('A2','SYN'))))
    monkeypatch.setattr(reader,setting,1)
    with pytest.raises(XLSXLimitError):
        read(tmp_path,source,skip=True)


def test_cancel_propagates_and_shared_database_is_removed(tmp_path,monkeypatch):
    source=book(tmp_path/'contract.xlsx',sheet(row(1,cell('A1','0','s'))+row(2,cell('A2','1','s'))),shared=['code','SYN'])
    locations=[]
    original=reader._Strings
    class Capture(original):
        def __init__(self,path):
            locations.append(Path(path).parent)
            super().__init__(path)
        def load(self,package,part):
            super().load(package,part)
            raise Cancelled('Отмена')
    monkeypatch.setattr(reader,'_Strings',Capture)
    with pytest.raises(Cancelled):
        read(tmp_path,source,skip=True)
    assert locations and all(not p.exists() for p in locations)


def test_shared_cache_is_bounded_by_bytes(tmp_path,monkeypatch):
    monkeypatch.setattr(reader,'MAX_CACHE',1024)
    strings=_Strings(tmp_path/'strings.sqlite')
    try:
        for index in range(100):
            strings.add('SYN_'+str(index)+'x'*100)
        for index in range(100):
            assert strings.get(str(index)).startswith('SYN_')
            assert strings.bytes<=1024
        assert len(strings.cache)<3
    finally:
        strings.close()


def test_corrupt_package_safe_error(tmp_path):
    source=tmp_path/'contract.xlsx'
    source.write_bytes(b'SYN_BROKEN_ZIP')
    with pytest.raises(AnonymizerError) as error:
        read(tmp_path,source)
    assert 'SYN_BROKEN' not in str(error.value)


@pytest.mark.parametrize('content', ['<v>1</v><v>2</v>', '<t>PRIVATE</t>'])
def test_malformed_cell_structure_rejected(tmp_path,content):
    source=book(tmp_path/'contract.xlsx',sheet(row(1,cell('A1','code'),cell('B1','qty'))+row(2,cell('A2','SYN'),'<c r="B2">'+content+'</c>')))
    with pytest.raises(AnonymizerError):
        read(tmp_path,source)


def test_trailing_formatting_only_cells_do_not_expand_table(tmp_path):
    source=book(tmp_path/'contract.xlsx',sheet(row(1,cell('A1','code'),cell('Z1'))+row(2,cell('A2','SYN'),cell('Z2'))))
    rows,count,_=read(tmp_path,source)
    assert count==1 and rows[0]['code']=='SYN'


@pytest.mark.parametrize('workbook_protection', [False,True])
def test_protected_books_and_sheets_rejected(tmp_path,workbook_protection):
    extra={}
    tail='<sheetProtection sheet="1"/>'
    if workbook_protection:
        extra['xl/workbook.xml']=f'<workbook xmlns="{NS}" xmlns:r="{REL}"><workbookProtection lockStructure="1"/><sheets><sheet name="SYN" sheetId="1" r:id="sheet1"/></sheets></workbook>'
        tail=''
    source=book(tmp_path/'contract.xlsx',sheet(row(1,cell('A1','code'))+row(2,cell('A2','SYN')),tail),extra=extra)
    rows,count,dropped=read(tmp_path,source,skip=True)
    assert not rows and count==0 and dropped


def test_fragmented_text_has_separate_byte_budget(tmp_path,monkeypatch):
    monkeypatch.setattr(reader,'MAX_FIELD',256)
    rich='<c r="A2" t="inlineStr"><is>'+''.join('<r><t>x</t></r>' for _ in range(100))+'</is></c>'
    source=book(tmp_path/'contract.xlsx',sheet(row(1,cell('A1','code'))+row(2,rich)))
    with pytest.raises(XLSXLimitError):
        read(tmp_path,source)


@pytest.mark.parametrize('comment_size', [0,65535])
def test_zip64_locator_cannot_bypass_preflight(tmp_path,monkeypatch,comment_size):
    source=book(tmp_path/'contract.xlsx',sheet(row(1,cell('A1','code'))+row(2,cell('A2','SYN'))))
    original=source.read_bytes()
    offset=original.rfind(b'PK\x05\x06')
    eocd=struct.unpack_from('<4s4H2IH',original,offset)
    count, directory_size, directory_offset=eocd[4:7]
    zip64=struct.pack('<4sQ2H2I4Q',b'PK\x06\x06',44,45,45,0,0,count,count,directory_size,directory_offset)
    locator=struct.pack('<4sIQI',b'PK\x06\x07',0,offset,1)
    # All ordinary fields are zero rather than sentinels; stdlib still follows
    # the locator and reads the real ZIP64 directory.
    ordinary=struct.pack('<4s4H2IH',b'PK\x05\x06',0,0,0,0,0,0,comment_size)
    source.write_bytes(original[:offset]+zip64+locator+ordinary+b'x'*comment_size)
    with zipfile.ZipFile(source) as parsed:
        assert len(parsed.infolist())==count>0
    def unbounded_constructor(*args,**kwargs):
        pytest.fail('ZipFile must not be constructed for ZIP64 catalogues')
    monkeypatch.setattr(reader.zipfile,'ZipFile',unbounded_constructor)
    with pytest.raises(AnonymizerError,match='ZIP64'):
        read(tmp_path,source)


@pytest.mark.parametrize('serial',['1e1000000','-1e1000000','1e-1000000','-1e-1000000'])
def test_extreme_date_exponents_are_safe_and_skip_entire_table(tmp_path,serial):
    source=book(tmp_path/'contract.xlsx',sheet(
        row(1,cell('A1','code'),cell('B1','day'))
        +row(2,cell('A2','SYN_VALID'),cell('B2','1','n'))
        +row(3,cell('A3','SYN_INVALID'),cell('B3',serial,'n'))))
    with pytest.raises(AnonymizerError):
        read(tmp_path,source)
    rows,count,dropped=read(tmp_path,source,skip=True)
    assert rows==[] and count==0 and dropped



def test_loaded_shared_strings_keep_one_bounded_read_snapshot(tmp_path,monkeypatch):
    monkeypatch.setattr(reader,'MAX_CACHE',0)
    source=book(tmp_path/'contract.xlsx',sheet(row(1,cell('A1','0','s'))),shared=['code','SYN_A','SYN_B'])
    package=reader._Package(source,lambda:None)
    strings=_Strings(tmp_path/'lookup.sqlite')
    try:
        strings.load(package,'xl/sharedStrings.xml')
        assert strings.db.in_transaction
        for index in range(30):
            assert strings.get(str(index%3))==['code','SYN_A','SYN_B'][index%3]
            assert strings.db.in_transaction
        assert strings.bytes==0 and not strings.cache
        assert strings.db.execute('PRAGMA cache_size').fetchone()[0]==-2048
    finally:
        strings.close()
        package.close()
