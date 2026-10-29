"""Synthetic generic XLSX import keeps native scalar types without model lookup."""
from datetime import date, datetime, time
import zipfile

import pytest
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from anonymizer.errors import AnonymizerError
from anonymizer.metadata import Model
from anonymizer.store import Store
from anonymizer.xlsx_reader import parse_xlsx
from test_xlsx_reader import NS, REL, PKG, book, cell, row, sheet


def read(source, tmp_path, *, skip=False):
    with Store(tmp_path / "spool.sqlite", Model(), bypass_model=True) as store:
        parse_xlsx(source, store, skip_error_tables=skip)
        if not store.tables:
            return {}, [], store.dropped_tables
        name, table = next(iter(store.tables.items()))
        kinds = {field["name"]: field["type"] for field in table["fields"]}
        values = [dict(zip(store.columns(name), record)) for record in store.rows(table["id"])]
        return kinds, values, store.dropped_tables


def styled_book(path, data, formats, prefix=""):
    styles = (f'<styleSheet xmlns="{NS}"><numFmts>'
              + ''.join(f'<numFmt numFmtId="{164+i}" formatCode="{fmt}"/>' for i, fmt in enumerate(formats))
              + '</numFmts><cellXfs><xf numFmtId="0"/>'
              + ''.join(f'<xf numFmtId="{164+i}"/>' for i in range(len(formats)))
              + '</cellXfs></styleSheet>')
    relationships = (f'<Relationships xmlns="{PKG}">'
                     f'<Relationship Id="sheet1" Type="{REL}/worksheet" Target="worksheets/sheet1.xml"/>'
                     f'<Relationship Id="styles" Type="{REL}/styles" Target="styles.xml"/>'
                     '</Relationships>')
    return book(path, f'<worksheet xmlns="{NS}">{prefix}<sheetData>{data}</sheetData></worksheet>', extra={
        "xl/styles.xml": styles, "xl/_rels/workbook.xml.rels": relationships})


def test_unknown_table_and_native_text_number_boolean(tmp_path):
    header = row(1, cell('A1', 'code'), cell('B1', 'amount'), cell('C1', 'flag'), cell('D1', 'note'))
    data = row(2, cell('A2', '123'), cell('B2', '12345678901234567890.123456789', 'n'),
               cell('C2', '1', 'b'), cell('D2', ''))
    data += row(3, cell('A3', '001'), cell('B3'), cell('C3', '0', 'b'), cell('D3', r'\N'))
    source = book(tmp_path / 'outside_model.xlsx', sheet(header + data))
    kinds, rows, dropped = read(source, tmp_path)
    assert not dropped
    assert kinds == {'code': 'TEXT', 'amount': 'NUMBER', 'flag': 'BOOLEAN', 'note': 'TEXT'}
    assert rows == [{'code': '123', 'amount': '12345678901234567890.123456789', 'flag': 'true', 'note': ''},
                    {'code': '001', 'amount': None, 'flag': 'false', 'note': None}]


def test_known_table_missing_model_business_keys_and_extra_column(tmp_path):
    source = book(tmp_path / 'sku.xlsx', sheet(row(1, cell('A1', 'extra_text'))
                                             + row(2, cell('A2', 'SYN_SECRET'))))
    kinds, rows, dropped = read(source, tmp_path)
    assert kinds == {'extra_text': 'TEXT'} and rows == [{'extra_text': 'SYN_SECRET'}]
    assert not dropped


@pytest.mark.parametrize('epoch1904', [False, True])
def test_actual_xlsx_style_dates_datetimes_and_times(tmp_path, epoch1904):
    source = tmp_path / 'outside_model.xlsx'
    workbook = Workbook()
    if epoch1904:
        workbook.epoch = CALENDAR_MAC_1904
    workbook.active.append(['day', 'stamp', 'clock', 'code'])
    workbook.active.append([date(2026, 10, 29), datetime(2026, 10, 29, 12), time(6), '001'])
    workbook.save(source)
    workbook.close()
    kinds, rows, _ = read(source, tmp_path)
    assert kinds == {'day': 'DATE', 'stamp': 'DATE_TIME', 'clock': 'TIME', 'code': 'TEXT'}
    assert rows == [{'day': '2026-10-29', 'stamp': '2026-10-29T12:00:00', 'clock': '06:00:00', 'code': '001'}]


def test_iso_typed_dates_and_text_dates_are_distinct(tmp_path):
    header = row(1, *[cell(f'{chr(65+i)}1', label) for i, label in enumerate(['day','stamp','clock','literal'])])
    data = row(2, cell('A2', '2026-10-29', 'd'), cell('B2', '2026-10-29T12:00:00', 'd'),
               cell('C2', '06:00:00', 'd'), cell('D2', '2026-10-29'))
    source = book(tmp_path / 'outside_model.xlsx', sheet(header + data))
    kinds, rows, _ = read(source, tmp_path)
    assert kinds == {'day':'DATE', 'stamp':'DATE_TIME', 'clock':'TIME', 'literal':'TEXT'}
    assert rows[0]['literal'] == '2026-10-29'


@pytest.mark.parametrize('inherited', ['cell', 'row', 'column'])
def test_custom_uppercase_date_format_and_inheritance(tmp_path, inherited):
    header = row(1, cell('A1', 'day'))
    value = cell('A2', '61', 'n')
    if inherited == 'cell':
        value = value.replace('r="A2"', 'r="A2" s="1"')
    data = header + row(2, value, attrs='s="1"' if inherited == 'row' else '')
    prefix = '<cols><col min="1" max="1" style="1"/></cols>' if inherited == 'column' else ''
    source = styled_book(tmp_path / 'outside_model.xlsx', data, ['YYYY-MM-DD'], prefix)
    kinds, rows, _ = read(source, tmp_path)
    assert kinds == {'day':'DATE'} and rows == [{'day':'1900-03-01'}]


def test_time_format_literal_does_not_turn_time_into_datetime(tmp_path):
    header = row(1, cell('A1', 'clock'))
    data = header + row(2, cell('A2', '.25', 'n').replace('r="A2"', 'r="A2" s="1"'))
    source = styled_book(tmp_path / 'outside_model.xlsx', data, ['hh:mm:ss &quot;day&quot;'])
    kinds, rows, _ = read(source, tmp_path)
    assert kinds == {'clock':'TIME'} and rows == [{'clock':'06:00:00'}]


@pytest.mark.parametrize('fmt', ['m', 'mm', 'mmm', 'mmmm', 'mmmmm', '&quot;h&quot; mmmm', r'mmmm \h'])
def test_month_only_formats_are_calendar_dates(tmp_path, fmt):
    header = row(1, cell('A1', 'day'))
    data = header + row(2, cell('A2', '61', 'n').replace('r="A2"', 'r="A2" s="1"'))
    source = styled_book(tmp_path / 'outside_model.xlsx', data, [fmt])
    kinds, rows, _ = read(source, tmp_path)
    assert kinds == {'day':'DATE'} and rows == [{'day':'1900-03-01'}]


@pytest.mark.parametrize('fmt', ['mm:ss', 'h:mm', 'hh:mm', 'hh:mm:ss', 'AM/PM', 'am/pm', 'A/P'])
def test_minute_formats_remain_time_values(tmp_path, fmt):
    header = row(1, cell('A1', 'clock'))
    data = header + row(2, cell('A2', '.25', 'n').replace('r="A2"', 'r="A2" s="1"'))
    source = styled_book(tmp_path / 'outside_model.xlsx', data, [fmt])
    kinds, rows, _ = read(source, tmp_path)
    assert kinds == {'clock':'TIME'} and rows == [{'clock':'06:00:00'}]


@pytest.mark.parametrize('fmt,serial', [('[h]:mm:ss','1.5'),('yyyy-mm-dd','60'),('hh:mm:ss','.000000000000001')])
def test_unsupported_duration_phantom_date_and_loss_are_table_errors(tmp_path, fmt, serial):
    header = row(1, cell('A1', 'value'))
    data = header + row(2, cell('A2', serial, 'n').replace('r="A2"', 'r="A2" s="1"'))
    source = styled_book(tmp_path / 'outside_model.xlsx', data, [fmt])
    kinds, rows, dropped = read(source, tmp_path, skip=True)
    assert not kinds and not rows and len(dropped) == 1


def test_mixed_native_types_fail_without_disclosing_value(tmp_path):
    data = row(1, cell('A1', 'mixed')) + row(2, cell('A2','1','n')) + row(3,cell('A3','SYN_PRIVATE_TEXT'))
    source = book(tmp_path / 'outside_model.xlsx', sheet(data))
    with pytest.raises(AnonymizerError) as error:
        read(source, tmp_path)
    assert 'SYN_PRIVATE_TEXT' not in str(error.value)


def test_formula_rejection_still_drops_all_generic_rows(tmp_path):
    data = row(1, cell('A1', 'note')) + row(2, cell('A2','SYN_PRIVATE'))
    data += '<row r="3"><c r="A3"><f>1+1</f><v>2</v></c></row>'
    source = book(tmp_path / 'outside_model.xlsx', sheet(data))
    kinds, rows, dropped = read(source, tmp_path, skip=True)
    assert not kinds and not rows and len(dropped) == 1
